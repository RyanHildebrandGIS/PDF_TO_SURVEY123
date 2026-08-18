import os

import openpyxl


def build_sample_template(path: str) -> None:
    if os.path.exists(path):
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)

    wb = openpyxl.Workbook()
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label", "required", "appearance"])
    survey.append(["geopoint", "site_location", "Site location", "yes", ""])
    survey.append(["select_one yes_no", "access_ok", "Site access OK?", "no", ""])
    survey.append(["begin group", "grp_form_content", "<b>[Form-Specific Content]</b>", "", ""])
    survey.append(["end group", "", "", "", ""])
    survey.append(["text", "inspector_notes", "Inspector notes", "no", "multiline"])

    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    for name, label in (("good", "Good"), ("fair", "Fair"), ("poor", "Poor")):
        choices.append(["condition", name, label])
    for name, label in (("yes", "Yes"), ("no", "No")):
        choices.append(["yes_no", name, label])

    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id"])
    settings.append(["Standard Inspection v4", "standard_inspection_v4"])

    wb.save(path)
